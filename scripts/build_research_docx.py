from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "RESEARCH_PAPER.md"
OUTPUT = ROOT / "NetVision_AI_Research_Paper.docx"


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(10)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_width(table, widths: list[int]) -> None:
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.tcW
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr_text)
    run._r.append(fld_end)


def next_numbering_id(numbering, tag: str) -> int:
    values = []
    for element in numbering.findall(qn(tag)):
        attr = "w:abstractNumId" if tag.endswith("abstractNum") else "w:numId"
        value = element.get(qn(attr))
        if value is not None:
            values.append(int(value))
    return (max(values) + 1) if values else 1


def create_numbered_list(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    abstract_id = next_numbering_id(numbering, "w:abstractNum")
    num_id = next_numbering_id(numbering, "w:num")

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")

    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)

    for node in (start, num_fmt, lvl_text, lvl_jc, p_pr):
        lvl.append(node)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)
    p_pr.append(num_pr)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.167


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor(31, 77, 120)
    title.paragraph_format.space_after = Pt(8)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(11)
    subtitle.font.italic = True
    subtitle.font.color.rgb = RGBColor(89, 89, 89)
    subtitle.paragraph_format.space_after = Pt(16)

    h1 = styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(46, 116, 181)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)

    h2 = styles["Heading 2"]
    h2.font.name = "Calibri"
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(46, 116, 181)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)

    h3 = styles["Heading 3"]
    h3.font.name = "Calibri"
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(31, 77, 120)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    header_para = header.paragraphs[0]
    header_para.text = "NetVision AI Research Paper"
    header_para.style = doc.styles["Header"]
    header_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header_para.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(89, 89, 89)

    footer = section.footer
    footer_para = footer.paragraphs[0]
    add_page_number(footer_para)
    for run in footer_para.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(89, 89, 89)


def add_metadata_table(doc: Document) -> None:
    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    rows = [
        ("Document type", "IEEE-style university software engineering manuscript"),
        ("System", "NetVision AI desktop network monitoring and management application"),
        ("Source basis", "Final repository implementation only"),
        ("Evaluation note", "Code inspection, build verification, endpoint smoke testing, and runtime validation; no controlled benchmarks"),
    ]
    for row_idx, (label, value) in enumerate(rows):
        set_cell_text(table.cell(row_idx, 0), label, bold=True)
        set_cell_text(table.cell(row_idx, 1), value)
        shade_cell(table.cell(row_idx, 0), "F2F4F7")
    set_table_width(table, [2500, 6860])


def add_paragraph_with_inline_bold(doc: Document, text: str, style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)
    return paragraph


def clean_text(text: str) -> str:
    return (
        text.replace("**", "")
        .replace(" - ", " - ")
        .replace("–", "-")
        .replace("—", "-")
    )


def build_docx() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    configure_styles(doc)
    add_header_footer(doc)

    lines = markdown.splitlines()
    title = lines[0].lstrip("# ").strip()
    title_para = doc.add_paragraph(title, style="Title")
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph(
        "Complete academic manuscript generated from the verified final implementation",
        style="Subtitle",
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_metadata_table(doc)

    in_references = False
    pending_blank = False
    current_num_id: int | None = None
    for raw in lines[1:]:
        line = raw.strip()
        if not line:
            pending_blank = True
            current_num_id = None
            continue

        if line.startswith("## References"):
            doc.add_page_break()
            doc.add_heading("References", level=1)
            in_references = True
            current_num_id = None
            continue

        if line.startswith("## "):
            heading = line[3:].strip()
            if heading in {"Abstract", "Keywords"}:
                doc.add_heading(heading, level=1)
            else:
                doc.add_heading(heading, level=1)
            pending_blank = False
            current_num_id = None
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            pending_blank = False
            current_num_id = None
            continue

        if in_references and re.match(r"^\[\d+\]", line):
            p = doc.add_paragraph(style="Normal")
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.space_after = Pt(4)
            p.add_run(line)
            pending_blank = False
            current_num_id = None
            continue

        if re.match(r"^\d+\.\s+", line):
            item = re.sub(r"^\d+\.\s+", "", line)
            if current_num_id is None:
                current_num_id = create_numbered_list(doc)
            paragraph = add_paragraph_with_inline_bold(doc, item)
            apply_numbering(paragraph, current_num_id)
            pending_blank = False
            continue

        if line.startswith("- "):
            add_paragraph_with_inline_bold(doc, line[2:].strip(), style="List Bullet")
            pending_blank = False
            current_num_id = None
            continue

        if line.startswith("**") and line.endswith("**") and len(line) < 120:
            p = doc.add_paragraph()
            run = p.add_run(line[2:-2])
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(11)
            pending_blank = False
            current_num_id = None
            continue

        paragraph = add_paragraph_with_inline_bold(doc, line)
        if pending_blank:
            paragraph.paragraph_format.space_before = Pt(0)
        pending_blank = False
        current_num_id = None

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_docx()
    print(OUTPUT)
